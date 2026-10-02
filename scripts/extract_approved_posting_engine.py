"""Extract reference data from the approved YIFRE posting workbook."""
from datetime import date, datetime
import json
from pathlib import Path
import sys

import openpyxl


def rows(sheet, first_data_row=6):
    header_values = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    result = []
    for values in sheet.iter_rows(min_row=first_data_row, values_only=True):
        row = {headers[index]: value for index, value in enumerate(values) if headers[index]}
        if not any(value is not None for value in row.values()):
            continue
        result.append(row)
    return result


def serialise(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise TypeError(f"Unsupported workbook value: {type(value).__name__}")


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract_approved_posting_engine.py WORKBOOK OUTPUT_JSON")
    workbook = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    names = {
        "accounts": "COA",
        "postingControls": "Posting Controls",
        "postingKeys": "Posting Keys",
        "postingGroups": "Posting Groups",
        "accountMaps": "Account Map",
        "costCentres": "Cost Centres",
    }
    data = {
        "source": Path(sys.argv[1]).name,
        "sheets": {key: rows(workbook[sheet]) for key, sheet in names.items()},
    }
    Path(sys.argv[2]).write_text(json.dumps(data, indent=2, default=serialise) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
